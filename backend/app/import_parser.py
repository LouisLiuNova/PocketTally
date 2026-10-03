"""个人账单安全解析；不写数据库，不执行公式。"""

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError as XMLParseError
from zipfile import BadZipFile, ZipFile
from zlib import error as ZlibError

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.time_utils import SHANGHAI, require_aware_datetime

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_MINOR = 9_000_000_000_000_000


class ParseError(ValueError):
    """仅包含不带账单正文的公开错误。"""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass
class ParsedRow:
    row_number: int
    source_id: str
    original: dict
    state: str
    reason: str | None
    transaction_type: str | None
    amount_minor: int
    occurred_at: datetime | None
    description: str
    payment_key: str


@dataclass
class ParsedBill:
    provider: str
    rows: list[ParsedRow]
    summary: dict


def clean(value) -> str:
    return str(value if value is not None else "").strip().lstrip("\ufeff").strip()


def amount_minor(value: str) -> int:
    try:
        amount = Decimal(re.sub(r"[¥￥,\s]", "", value))
        if not amount.is_finite() or amount <= 0 or amount.as_tuple().exponent < -2:
            raise ValueError()
        minor = int(amount * 100)
        if minor > MAX_MINOR:
            raise ValueError()
        return minor
    except InvalidOperation, ValueError, OverflowError:
        raise ValueError("金额格式不合法或超出整数分范围") from None


def read_rows(data: bytes, filename: str) -> list[tuple[int, list[str]]]:
    suffix = Path(filename).suffix.lower()
    if not data or len(data) > MAX_FILE_BYTES:
        raise ParseError("import_file_limit", "文件为空或超过 10 MiB")
    if suffix == ".csv":
        if data.startswith(b"PK") or b"\x00" in data:
            raise ParseError("import_invalid_file", "文件内容不是受支持的 CSV")
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                content = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ParseError("import_encoding", "CSV 编码无法识别")
        try:
            reader = csv.reader(io.StringIO(content, newline=""), strict=True)
            result = []
            for row in reader:
                result.append((reader.line_num, [clean(cell) for cell in row]))
                if len(result) > MAX_ROWS + 200:
                    raise ParseError("import_row_limit", "交易数据行超过 10,000 行")
            return result
        except csv.Error:
            raise ParseError("import_invalid_file", "CSV 格式损坏") from None
    if suffix != ".xlsx" or not data.startswith(b"PK"):
        raise ParseError("import_unsupported_file", "仅支持支付宝 CSV、微信 CSV/XLSX")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES
                or len(entries) > 1000
            ):
                raise ParseError("import_file_limit", "工作簿解压后超过安全限制")
            if (
                any(item.flag_bits & 1 for item in entries)
                or "xl/vbaProject.bin" in archive.namelist()
            ):
                raise ParseError("import_invalid_file", "不支持加密或包含宏的工作簿")
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) != 1:
                raise ParseError("import_invalid_file", "个人账单应包含一个工作表")
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            result = []
            for number, row in enumerate(sheet.iter_rows(), 1):
                if any(cell.data_type == "f" for cell in row):
                    raise ParseError(
                        "import_formula", "账单包含公式单元格，请重新导出原始账单"
                    )
                if len(row) > 100:
                    raise ParseError("import_file_limit", "工作簿列数超过安全限制")
                result.append((number, [clean(cell.value) for cell in row]))
                if len(result) > MAX_ROWS + 200:
                    raise ParseError("import_row_limit", "交易数据行超过 10,000 行")
            return result
        finally:
            workbook.close()
    except ParseError:
        raise
    except (
        BadZipFile,
        XMLParseError,
        DefusedXmlException,
        InvalidFileException,
        ValueError,
        TypeError,
        KeyError,
        IndexError,
        OSError,
        EOFError,
        ZlibError,
    ):
        raise ParseError("import_invalid_file", "工作簿损坏或格式不受支持") from None


def parse_bill(data: bytes, filename: str) -> ParsedBill:
    raw = read_rows(data, filename)
    header_index = None
    provider = None
    for i, (_, values) in enumerate(raw[:200]):
        if "交易时间" in values and "收/支" in values:
            if "交易订单号" in values and "交易状态" in values:
                provider = "alipay"
            elif "交易单号" in values and "当前状态" in values:
                provider = "wechat"
            if provider:
                header_index = i
                break
    if header_index is None or (
        provider == "alipay" and Path(filename).suffix.lower() != ".csv"
    ):
        raise ParseError("import_unknown_format", "未识别为受支持的个人账单表头")
    header = raw[header_index][1]
    required = {"交易时间", "交易对方", "收/支"}
    required |= (
        {"交易订单号", "交易状态", "金额", "收/付款方式", "商品说明", "交易分类"}
        if provider == "alipay"
        else {"交易单号", "当前状态", "金额(元)", "支付方式", "商品", "交易类型"}
    )
    if not required.issubset(header) or len([h for h in header if h]) != len(
        {h for h in header if h}
    ):
        raise ParseError("import_unknown_format", "账单缺少必要列或存在重复表头")
    candidates = []
    seen: dict[str, ParsedRow] = {}
    for number, values in raw[header_index + 1 :]:
        if not any(values):
            continue
        if len(candidates) >= MAX_ROWS:
            raise ParseError("import_row_limit", "交易数据行超过 10,000 行")
        source = dict(zip(header, values))
        get = lambda a, b, source=source: source.get(
            a if provider == "alipay" else b, ""
        )
        source_id = get("交易订单号", "交易单号").lstrip("'\t")
        kind = get("交易分类", "交易类型")
        status = get("交易状态", "当前状态")
        direction = source.get("收/支", "")
        product = get("商品说明", "商品")
        payment = get("收/付款方式", "支付方式")
        original = {
            "sourceId": source_id,
            "merchantOrderId": get("商家订单号", "商户单号"),
            "counterparty": source.get("交易对方", ""),
            "product": product,
            "paymentMethod": payment,
            "platformCategory": kind,
            "platformStatus": status,
            "direction": direction,
            "amount": get("金额", "金额(元)"),
            "occurredAt": source.get("交易时间", ""),
            "rowNumbers": [number],
        }
        state, reason = "candidate", None
        transaction_type = {"收入": "income", "支出": "expense"}.get(direction)
        amount, occurred = 0, None
        try:
            amount = amount_minor(original["amount"])
            occurred = require_aware_datetime(
                datetime.strptime(original["occurredAt"], "%Y-%m-%d %H:%M:%S").replace(
                    tzinfo=SHANGHAI
                )
            ).replace(tzinfo=None)
        except ValueError as error:
            state, reason = (
                "unsupported",
                str(error) if amount == 0 else "交易时间格式不合法",
            )
        # 退款行与原支出上的已退款状态必须区分，不能把原支出误认为退款收入。
        is_refund = (
            kind == "退款"
            or kind.endswith("-退款")
            or status == "退款成功"
            or (direction == "收入" and "退款" in status)
        )
        if is_refund:
            state, reason = (
                "refund_pending",
                "首版不导入退款，请单独核对并使用支出退款功能",
            )
        elif direction in {"不计收支", "/"} or kind in {
            "投资理财",
            "信用借还",
            "账户存取",
            "零钱提现",
            "充值",
            "提现",
        }:
            state, reason = "skipped", "不计收支或内部资金变动"
        elif status in {"交易关闭", "支付失败", "交易失败", "已关闭", "已撤销"}:
            state, reason = "skipped", "交易关闭、失败或撤销"
        elif re.fullmatch(r"[¥￥]?0(?:\.0+)?", original["amount"]):
            state, reason = "skipped", "零金额交易，不产生资金变动"
        elif transaction_type is None:
            state, reason = "unsupported", "无法识别收支方向"
        elif not (
            status
            in {
                "交易成功",
                "代付成功",
                "支付成功",
                "转账成功",
                "对方已收钱",
                "已转账",
                "已全额退款",
                "对方已退还",
            }
            or status.startswith("已退款")
        ):
            state, reason = "unsupported", "交易状态尚不支持"
        elif not source_id:
            state, reason = "unsupported", "缺少平台交易单号"
        if any(value.startswith("=") for value in values):
            state, reason = "unsupported", "记录包含公式形式的字段"
        row = ParsedRow(
            number,
            source_id,
            original,
            state,
            reason,
            transaction_type,
            amount,
            occurred,
            " · ".join(filter(None, [source.get("交易对方", ""), product]))[:2000],
            payment,
        )
        if state == "candidate" and source_id and source_id in seen:
            previous = seen[source_id]
            comparable = lambda item: {
                k: v for k, v in item.original.items() if k != "rowNumbers"
            }
            if comparable(previous) == comparable(row):
                previous.original["rowNumbers"].append(number)
                row.state, row.reason = "skipped", "相同来源单号的重复明细已折叠"
            else:
                previous.state = row.state = "unsupported"
                previous.reason = row.reason = "同一来源单号的明细冲突，不能安全合并"
        elif state == "candidate" and source_id:
            seen[source_id] = row
        candidates.append(row)
    if not candidates:
        raise ParseError("import_empty", "账单没有交易数据行")
    counts = dict(Counter(r.state for r in candidates))
    totals = {
        direction: sum(
            r.amount_minor for r in candidates if r.original["direction"] == direction
        )
        for direction in ("收入", "支出", "不计收支", "/")
    }
    dates = [r.occurred_at.isoformat() + "Z" for r in candidates if r.occurred_at]
    return ParsedBill(
        provider,
        candidates,
        {
            "inputRows": len(candidates),
            "parseCounts": counts,
            "sourceAmountMinor": totals,
            "startAt": min(dates) if dates else None,
            "endAt": max(dates) if dates else None,
        },
    )
